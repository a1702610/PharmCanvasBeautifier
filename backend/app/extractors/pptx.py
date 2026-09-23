import io

from pptx import Presentation
from pptx.shapes.group import GroupShape
from pptx.shapes.picture import Picture

from app.extractors.common import RawImage, SourceExtract, table_to_text


def _iter_shapes(shapes):
    for shape in shapes:
        if isinstance(shape, GroupShape):
            yield from _iter_shapes(shape.shapes)
        else:
            yield shape


def _para_text(para) -> str:
    pieces = []
    for run in para.runs:
        address = run.hyperlink.address if run.hyperlink is not None else None
        if address and run.text.strip():
            pieces.append(f"[{run.text}]({address})")
        else:
            pieces.append(run.text)
    return "".join(pieces).strip()


def extract_pptx(data: bytes, name: str) -> SourceExtract:
    prs = Presentation(io.BytesIO(data))
    out = SourceExtract()
    slides_text: list[str] = []
    for n, slide in enumerate(prs.slides, start=1):
        title_shape = slide.shapes.title
        title = title_shape.text_frame.text.strip() if title_shape is not None and title_shape.has_text_frame else ""
        lines = [f"--- Slide {n}: {title} ---" if title else f"--- Slide {n} ---"]
        for shape in _iter_shapes(slide.shapes):
            if title_shape is not None and shape.shape_id == title_shape.shape_id:
                continue
            if isinstance(shape, Picture):
                try:
                    blob = shape.image.blob
                except Exception:
                    continue
                lines.append(out.add_image(RawImage(source=name, location=f"slide {n}", data=blob)))
            elif getattr(shape, "has_table", False) and shape.has_table:
                text = table_to_text([[c.text for c in row.cells] for row in shape.table.rows])
                if text:
                    lines.append(text)
            elif shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    text = _para_text(para)
                    if text:
                        lines.append("  " * para.level + "- " + text)
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                lines.append("Speaker notes: " + notes)
        slides_text.append("\n".join(lines))
    out.text = "\n\n".join(slides_text)
    return out
