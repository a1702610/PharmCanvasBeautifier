import io
import os

from PIL import Image


def noise_png(w: int = 200, h: int = 200) -> bytes:
    """Random-noise PNG: incompressible, so it passes the 3 KB size filter."""
    img = Image.frombytes("RGB", (w, h), os.urandom(w * h * 3))
    buf = io.BytesIO()
    img.save(buf, "PNG")
    return buf.getvalue()


def noise_bmp(w: int = 200, h: int = 200) -> bytes:
    img = Image.frombytes("RGB", (w, h), os.urandom(w * h * 3))
    buf = io.BytesIO()
    img.save(buf, "BMP")
    return buf.getvalue()


def make_docx(png: bytes) -> bytes:
    from docx import Document
    from docx.shared import Inches

    doc = Document()
    doc.add_heading("Overview", level=1)
    doc.add_paragraph("Chronic pain persists beyond 3 months.")
    doc.add_paragraph("Sensitisation", style="List Bullet")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Class"
    table.cell(0, 1).text = "Agents"
    table.cell(1, 0).text = "TCA"
    table.cell(1, 1).text = "amitriptyline"
    doc.add_picture(io.BytesIO(png), width=Inches(2))
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def make_pptx(png: bytes, logo: bytes | None = None, slides: int = 1) -> bytes:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    for n in range(slides):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        if logo is not None:
            slide.shapes.add_picture(io.BytesIO(logo), Inches(8), Inches(0.2), Inches(1))
        if n > 0:
            slide.shapes.title.text = f"Slide {n + 1}"
            continue
        slide.shapes.title.text = "Opioids"
        body = slide.placeholders[1].text_frame
        body.text = "Start low"
        sub = body.add_paragraph()
        sub.text = "Go slow"
        sub.level = 1
        linked = body.add_paragraph()
        run = linked.add_run()
        run.text = "TGA guidance"
        run.hyperlink.address = "https://www.tga.gov.au"
        slide.notes_slide.notes_text_frame.text = "Explain tolerance"
        slide.shapes.add_picture(io.BytesIO(png), Inches(1), Inches(1))
        table = slide.shapes.add_table(2, 2, Inches(1), Inches(5), Inches(4), Inches(1)).table
        table.cell(0, 0).text = "Class"
        table.cell(0, 1).text = "Agents"
        table.cell(1, 0).text = "Opioid"
        table.cell(1, 1).text = "morphine"
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()
